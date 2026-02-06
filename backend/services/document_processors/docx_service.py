from pathlib import Path
import logging
logger = logging.getLogger(__name__)


class DOCXService:
    """Service for DOCX file extraction and generation"""
    
    def extract_text(self, file_path: Path) -> str:
        """Extract text content from a DOCX file"""
        try:
            # Try to import python-docx, if not available, raise helpful error
            try:
                from docx import Document
            except ImportError:
                raise Exception(
                    "python-docx library is required for DOCX support. "
                    "Install it with: pip install python-docx"
                )
            
            doc = Document(file_path)
            text_content = []
            
            # Extract text from paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_content.append(paragraph.text)
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        text_content.append(' | '.join(row_text))
            
            return '\n\n'.join(text_content)
        except Exception as e:
            logger.error(f"Error reading DOCX file: {e}")
            raise Exception(f"Failed to extract text from DOCX file: {str(e)}")
    
    def create_file_from_text(self, text: str, output_path: Path):
        """Create a DOCX file from text content"""
        try:
            # Import python-docx
            try:
                from docx import Document
            except ImportError:
                raise Exception(
                    "python-docx library is required for DOCX support. "
                    "Install it with: pip install python-docx"
                )
            
            # Create a new Document
            doc = Document()
            
            # Split text into paragraphs (by double newlines)
            paragraphs = text.split('\n\n')
            
            for para_text in paragraphs:
                if para_text.strip():
                    # Remove single newlines within paragraphs and add as paragraph
                    cleaned_text = para_text.strip().replace('\n', ' ')
                    doc.add_paragraph(cleaned_text)
            
            # Save the document
            doc.save(str(output_path))
            logger.info(f"DOCX file created successfully at {output_path}")
        
        except Exception as e:
            logger.error(f"Error creating DOCX file: {e}")
            raise Exception(f"Failed to create DOCX file: {str(e)}")

